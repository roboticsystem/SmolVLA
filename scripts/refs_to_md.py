#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""refs_to_md.py —— 把「参考论文」目录里的 PDF / DOCX / PPTX / DOC 批量转成 Markdown。

为什么要这一步：参考论文是 PDF/DOCX 二进制，每次让模型"读一遍再转写"既慢又
不可复用（表格/公式在纯文本抽取里会整段丢失）。本脚本把每篇论文转成一份
**结构化 Markdown**（标题层级 / 表格转 md 表格 / 图片导出并内联引用），
落在 `参考论文/md/`，之后任何分析、引用、比对都直接读 md 即可。

用法:
    python scripts/refs_to_md.py                 # 转全部
    python scripts/refs_to_md.py --only 宋志光     # 只转名字含关键字的
    python scripts/refs_to_md.py --out 参考论文/md2
"""
from __future__ import annotations

import argparse
import os
import re
import struct
import sys
import zipfile
from collections import Counter

SRC_DIR = "参考论文"
OUT_DIR = os.path.join(SRC_DIR, "md")

# ---------------------------------------------------------------- 通用工具
CJK_NUM = r"第\s*[0-9一二三四五六七八九十]+\s*[章节]"
RE_CHAPTER = re.compile(CJK_NUM)
RE_NUMHEAD = re.compile(r"^\s*(\d+(?:\.\d+){0,3})\s*[、.\s]")


def _safe(name: str) -> str:
    return re.sub(r'[\\/:*?"<>|]+', "_", name).strip()


def _write(path: str, text: str):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        f.write(text)


# ---------------------------------------------------------------- PDF → MD
def pdf_to_md(pdf_path: str, out_md: str, media_dir: str) -> dict:
    import pymupdf

    doc = pymupdf.open(pdf_path)
    # ---- 统计正文字号（出现最多的字号即正文字号），用于标题判定
    sizes = Counter()
    for pno in range(doc.page_count):
        for blk in doc[pno].get_text("dict")["blocks"]:
            for ln in blk.get("lines", []):
                for sp in ln.get("spans", []):
                    t = sp.get("text", "").strip()
                    if t:
                        sizes[round(sp.get("size", 0), 1)] += len(t)
    body_size = sizes.most_common(1)[0][0] if sizes else 12.0

    md = []
    img_seq = 0
    tbl_seq = 0
    for pno in range(doc.page_count):
        page = doc[pno]
        md.append("\n<!-- page %d -->\n" % (pno + 1))
        # ---- 表格
        try:
            tabs = page.find_tables()
        except Exception:                                   # noqa: BLE001
            tabs = []
        table_boxes = []
        for tb in (tabs.tables if tabs else []):
            try:
                rows = tb.extract()
            except Exception:                               # noqa: BLE001
                continue
            if not rows or len(rows) < 2:
                continue
            tbl_seq += 1
            md.append("\n**表 %d（第 %d 页）**\n" % (tbl_seq, pno + 1))
            md.append(_rows_to_md(rows))
            md.append("\n")
            table_boxes.append(tuple(tb.bbox))
        # ---- 正文（按 y 排序，表格区域跳过）
        blocks = [b for b in page.get_text("dict")["blocks"]
                  if b.get("type", 0) == 0]
        blocks.sort(key=lambda b: (round(b["bbox"][1], 1), b["bbox"][0]))
        for blk in blocks:
            bb = tuple(blk["bbox"])
            cx, cy = (bb[0] + bb[2]) / 2.0, (bb[1] + bb[3]) / 2.0
            if any(tb[0] <= cx <= tb[2] and tb[1] <= cy <= tb[3]
                   for tb in table_boxes):
                continue
            for ln in blk.get("lines", []):
                spans = ln.get("spans", [])
                if not spans:
                    continue
                text = "".join(sp.get("text", "") for sp in spans).strip()
                if not text:
                    continue
                size = max(sp.get("size", 0) for sp in spans)
                bold = any("bold" in (sp.get("font", "") or "").lower()
                           or "黑体" in (sp.get("font", "") or "")
                           or "Bold" in (sp.get("font", "") or "")
                           for sp in spans)
                md.append(_pdf_line_to_md(text, size, body_size, bold) + "\n")
        # ---- 图片（降采样到 ≤1400px 宽，避免 md 媒体目录膨胀）
        for info in page.get_images(full=True):
            xref = info[0]
            try:
                pix = pymupdf.Pixmap(doc, xref)
                if pix.n - pix.alpha >= 4:
                    pix = pymupdf.Pixmap(pymupdf.csRGB, pix)
                img_seq += 1
                name = "p%03d_img%02d.png" % (pno + 1, img_seq)
                path = os.path.join(media_dir, name)
                if pix.width > 1400:
                    from io import BytesIO
                    from PIL import Image
                    im = Image.open(BytesIO(pix.tobytes("png"))).convert("RGB")
                    w = 1400
                    h = max(1, int(im.height * w / im.width))
                    im.resize((w, h)).save(path)
                else:
                    pix.save(path)
                md.append("\n![](%s)\n" % name)
            except Exception:                               # noqa: BLE001
                continue
    n_pages = doc.page_count
    doc.close()
    head = ("# %s\n\n> 由 `scripts/refs_to_md.py` 自动转换（PDF，正文字号 %.1fpt，"
            "共 %d 页 / %d 图 / %d 表）。\n\n"
            % (os.path.splitext(os.path.basename(pdf_path))[0], body_size,
               n_pages, img_seq, tbl_seq))
    _write(out_md, head + "".join(md))
    return {"pages": doc.page_count, "images": img_seq, "tables": tbl_seq}


def _pdf_line_to_md(text: str, size: float, body: float, bold: bool) -> str:
    if RE_CHAPTER.match(text) and size >= body - 0.5:
        return "# " + text
    if size >= body + 3.5:
        return "# " + text
    if size >= body + 1.8:
        return "## " + text
    m = RE_NUMHEAD.match(text)
    if m and (size >= body + 0.4 or bold):
        depth = m.group(1).count(".") + 1
        return "#" * min(depth + 1, 6) + " " + text
    if size >= body + 0.4 and bold and len(text) < 40:
        return "**%s**" % text
    return text


def _rows_to_md(rows) -> str:
    rows = [[("" if c is None else str(c).replace("\n", " ").strip())
             for c in r] for r in rows]
    if not rows:
        return ""
    ncol = max(len(r) for r in rows)
    rows = [r + [""] * (ncol - len(r)) for r in rows]
    out = ["| " + " | ".join(rows[0]) + " |",
           "| " + " | ".join(["---"] * ncol) + " |"]
    for r in rows[1:]:
        out.append("| " + " | ".join(r) + " |")
    return "\n".join(out)


# ---------------------------------------------------------------- DOCX → MD
def docx_to_md(docx_path: str, out_md: str, media_dir: str) -> dict:
    import docx
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(docx_path)
    body = d.element.body
    md = []
    img_seq = 0
    tbl_seq = 0

    def para_images(p_el, rels):
        """段落内嵌图片（w:drawing / w:pict）→ 导出并返回 md 引用列表"""
        nonlocal img_seq
        refs = []
        for blip in p_el.iter(qn("a:blip")):
            rid = blip.get(qn("r:embed")) or blip.get(qn("r:link"))
            if not rid or rid not in rels:
                continue
            try:
                blob = rels[rid].blob
            except Exception:                               # noqa: BLE001
                continue
            img_seq += 1
            name = "img%03d_%s.png" % (img_seq, _safe(rid))
            with open(os.path.join(media_dir, name), "wb") as f:
                f.write(blob)
            refs.append(name)
        return refs

    rels = d.part.related_parts
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            p = Paragraph(child, d)
            text = p.text.strip()
            style = (p.style.name or "") if p.style is not None else ""
            imgs = para_images(child, rels)
            if not text and not imgs:
                continue
            if re.search(r"heading\s*\d|^标题\s*\d", style, re.I):
                m = re.search(r"(\d+)", style)
                lvl = int(m.group(1)) if m else 1
                md.append("#" * min(lvl, 6) + " " + text + "\n")
            elif style.lower() in ("title",):
                md.append("# " + text + "\n")
            elif text:
                md.append(text + "\n")
            for im in imgs:
                md.append("\n![](%s)\n" % im)
        elif child.tag == qn("w:tbl"):
            t = Table(child, d)
            raw = [[c.text.strip() for c in r.cells] for r in t.rows]
            rows = [[c.replace("\n", " ") for c in r] for r in raw]
            tbl_seq += 1
            md.append("\n**表 %d**\n" % tbl_seq)
            md.append(_rows_to_md(rows) + "\n")
            # 中期检查表这类「一格装整段正文」的表：把长单元格原文按段落补出来，
            # 否则整段被压成一行，后续检索/引用都不好用
            for r in raw:
                for c in r:
                    if len(c) > 200:
                        md.append("\n**表 %d 长单元格原文**\n\n" % tbl_seq)
                        md.append(c + "\n")
    head = ("# %s\n\n> 由 `scripts/refs_to_md.py` 自动转换（DOCX：%d 段 / %d 图 / %d 表）。"
            "\n\n" % (os.path.splitext(os.path.basename(docx_path))[0],
                     len(d.paragraphs), img_seq, tbl_seq))
    _write(out_md, head + "\n".join(md))
    return {"images": img_seq, "tables": tbl_seq}


# ---------------------------------------------------------------- PPTX → MD
def _save_img(blob: bytes, ext: str, media_dir: str, name: str) -> bool:
    """图片落盘；宽 >1400px 的等比缩到 1400，避免 md 媒体目录过大。"""
    from io import BytesIO
    ext = (ext or "png").lower().replace("emf", "png").replace("wmf", "png")
    if ext not in ("png", "jpg", "jpeg", "gif", "bmp", "webp"):
        ext = "png"
    try:
        from PIL import Image
        im = Image.open(BytesIO(blob))
        if im.width > 1400:
            h = max(1, int(im.height * 1400 / im.width))
            im = im.convert("RGB").resize((1400, h))
        if ext == "jpg":
            ext = "jpeg"
        im.convert("RGB").save(os.path.join(media_dir, name), ext.upper()
                               if ext != "jpeg" else "JPEG", quality=90)
        return True
    except Exception:                                    # noqa: BLE001
        pass
    try:
        with open(os.path.join(media_dir, name), "wb") as f:
            f.write(blob)
        return True
    except Exception:                                    # noqa: BLE001
        return False


def pptx_to_md(pptx_path: str, out_md: str, media_dir: str) -> dict:
    """按「第 N 页」切分，页内按 上→下/左→右 顺序输出文本、表格、图片。"""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(pptx_path)
    md = []
    img_seq = 0
    tbl_seq = 0

    def walk(shapes):
        for sh in shapes:
            try:
                if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                    yield from walk(sh.shapes)
                    continue
            except Exception:                            # noqa: BLE001
                pass
            yield sh

    for sno, slide in enumerate(prs.slides, 1):
        md.append("\n## 第 %d 页\n" % sno)
        shapes = list(walk(slide.shapes))
        shapes.sort(key=lambda s: (int(getattr(s, "top", 0) or 0),
                                   int(getattr(s, "left", 0) or 0)))
        for sh in shapes:
            # ---- 表格
            if getattr(sh, "has_table", False):
                rows = [[c.text.replace("\n", " ").strip() for c in r.cells]
                        for r in sh.table.rows]
                tbl_seq += 1
                md.append("\n**表 %d（第 %d 页）**\n" % (tbl_seq, sno))
                md.append(_rows_to_md(rows) + "\n")
                continue
            # ---- 文本
            if getattr(sh, "has_text_frame", False):
                t = "\n".join(p.text for p in sh.text_frame.paragraphs).strip()
                if t:
                    md.append(t + "\n")
                continue
            # ---- 图片
            try:
                img = sh.image
            except Exception:                            # noqa: BLE001
                continue
            img_seq += 1
            name = "p%03d_img%02d.%s" % (sno, img_seq,
                                         (img.ext or "png").lower())
            name = re.sub(r"\.(emf|wmf)$", ".png", name)
            if _save_img(img.blob, img.ext, media_dir, name):
                md.append("\n![](%s)\n" % name)
        # ---- 备注页（中期汇报常把细节写在备注里）
        try:
            if slide.has_notes_slide:
                nt = slide.notes_slide.notes_text_frame.text.strip()
                if nt:
                    md.append("\n> 备注：%s\n" % nt.replace("\n", " / "))
        except Exception:                                # noqa: BLE001
            pass
    head = ("# %s\n\n> 由 `scripts/refs_to_md.py` 自动转换（PPTX：%d 页 / %d 图 / "
            "%d 表；表格按页顺序编号）。\n\n"
            % (os.path.splitext(os.path.basename(pptx_path))[0],
               len(prs.slides._sldIdLst), img_seq, tbl_seq))
    _write(out_md, head + "\n".join(md))
    return {"slides": len(prs.slides._sldIdLst), "images": img_seq,
            "tables": tbl_seq}


# ---------------------------------------------------------------- DOC(97-2003) → MD
def _doc_pieces(clx: bytes):
    """解析 Clx：跳过 Prc(0x01)，取 Pcdt(0x02) 里的片段表 [(cp0, cp1, fc)]。"""
    i = 0
    while i < len(clx):
        t = clx[i]
        if t == 1:                                       # Prc: 0x01+cbGrpprl(2)+grpprl
            i += 3 + struct.unpack_from("<h", clx, i + 1)[0]
        elif t == 2:                                     # Pcdt: 0x02+lcb(4)+PlcPcd
            lcb = struct.unpack_from("<I", clx, i + 1)[0]
            plc = clx[i + 5:i + 5 + lcb]
            n = (lcb - 4) // 12
            cps = struct.unpack_from("<%dI" % (n + 1), plc, 0)
            off = (n + 1) * 4
            return [(cps[k], cps[k + 1],
                     struct.unpack_from("<I", plc, off + k * 8 + 2)[0])
                    for k in range(n)]
        else:
            break
    return []


def doc_text(doc_path: str) -> str:
    """Word 97-2003 正文抽取：FIB → Clx 片段表 → 逐段解码（UTF-16 或 8 bit）。"""
    import olefile
    ole = olefile.OleFileIO(doc_path)
    try:
        word = ole.openstream("WordDocument").read()
        flags = struct.unpack_from("<H", word, 10)[0]
        tname = "1Table" if (flags >> 9) & 1 else "0Table"
        if not ole.exists(tname):
            tname = "0Table" if tname == "1Table" else "1Table"
        table = ole.openstream(tname).read()
        # FIB 基址 32B → csw → rgw97 → cslw → rglw97 → cbRgFcLcb → rgFcLcb97
        base = 32
        csw = struct.unpack_from("<H", word, base)[0]
        off = base + 2 + csw * 2
        cslw = struct.unpack_from("<H", word, off)[0]
        off += 2 + cslw * 4 + 2
        fc_clx, lcb_clx = struct.unpack_from("<II", word, off + 33 * 8)
        chunks = []
        for cp0, cp1, fc in _doc_pieces(table[fc_clx:fc_clx + lcb_clx]):
            n = cp1 - cp0
            if fc & 0x40000000:                          # 压缩片段：8 bit 本地码页
                s = (fc & 0x3FFFFFFF) // 2
                chunks.append(word[s:s + n].decode("cp1252", "replace"))
            else:
                s = fc & 0x3FFFFFFF
                chunks.append(word[s:s + n * 2].decode("utf-16-le", "replace"))
    finally:
        ole.close()
    txt = "".join(chunks)
    # Word 控制符：\r 段落、\x07 单元格结束、\x0b 软换行、\x0c 分页
    txt = (txt.replace("\x07", "\n").replace("\r", "\n")
              .replace("\x0b", "\n").replace("\x0c", "\n"))
    txt = "".join(ch for ch in txt
                  if ch in "\n\t" or (ord(ch) >= 32 and ord(ch) != 0xFFFD))
    return "\n".join(ln.rstrip() for ln in txt.split("\n"))


def doc_to_md(doc_path: str, out_md: str, media_dir: str) -> dict:
    text = doc_text(doc_path)
    md = []
    for ln in text.split("\n"):
        s = ln.strip()
        if not s:
            continue
        if re.match(r"^[一二三四五六七八九十]+、", s):        # 一、研究工作…
            md.append("\n## " + s + "\n")
        elif re.match(r"^（?[0-9]+）", s) and len(s) < 40:
            md.append("\n### " + s + "\n")                # 小节序号
        else:
            md.append(s + "\n")
    head = ("# %s\n\n> 由 `scripts/refs_to_md.py` 自动转换（DOC 97-2003，"
            "表格已拍平为段落；共 %d 字符）。\n\n"
            % (os.path.splitext(os.path.basename(doc_path))[0], len(text)))
    _write(out_md, head + "\n".join(md))
    return {"chars": len(text)}


# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", default=SRC_DIR)
    ap.add_argument("--out", default=OUT_DIR)
    ap.add_argument("--only", default=None, help="只转文件名含该关键字的")
    a = ap.parse_args()
    os.makedirs(a.out, exist_ok=True)
    files = [f for f in sorted(os.listdir(a.src))
             if f.lower().endswith((".pdf", ".docx", ".doc", ".pptx"))
             and not f.startswith("~$")]
    if a.only:
        files = [f for f in files if a.only in f]
    for f in files:
        stem = _safe(os.path.splitext(f)[0])
        md_path = os.path.join(a.out, stem + ".md")
        media = os.path.join(a.out, stem + "_media")
        os.makedirs(media, exist_ok=True)
        low = f.lower()
        try:
            if low.endswith(".pdf"):
                info = pdf_to_md(os.path.join(a.src, f), md_path, media)
            elif low.endswith(".docx"):
                info = docx_to_md(os.path.join(a.src, f), md_path, media)
            elif low.endswith(".pptx"):
                info = pptx_to_md(os.path.join(a.src, f), md_path, media)
            else:                                           # .doc（Word 97-2003）
                info = doc_to_md(os.path.join(a.src, f), md_path, media)
            print("[ok] %-60s -> %s  %s" % (f[:58], stem + ".md", info),
                  flush=True)
        except Exception as e:                              # noqa: BLE001
            print("[FAIL] %s: %s: %s" % (f, type(e).__name__, e), flush=True)
    # 索引
    idx = ["# 参考论文 Markdown 索引\n",
           "> 由 `scripts/refs_to_md.py` 生成；正文在同名 `.md`，图片在 `<论文名>_media/`。\n",
           "\n| 论文 | 文件 | 大小 | 标题数 | 表格数 | 图片数 |",
           "|---|---|---|---|---|---|"]
    for f in files:
        stem = _safe(os.path.splitext(f)[0])
        p = os.path.join(a.out, stem + ".md")
        if not os.path.exists(p):
            continue
        txt = open(p, encoding="utf-8").read()
        idx.append("| %s | [%s.md](%s.md) | %d KB | %d | %d | %d |" % (
            stem, stem, stem, os.path.getsize(p) // 1024,
            len(re.findall(r"^#{1,4} ", txt, re.M)),
            len(re.findall(r"^\*\*表", txt, re.M)),
            len(re.findall(r"!\[\]\(", txt))))
    _write(os.path.join(a.out, "INDEX.md"), "\n".join(idx) + "\n")
    print("[done] index -> %s" % os.path.join(a.out, "INDEX.md"))


if __name__ == "__main__":
    sys.exit(main())
